"""Read a student list from a Word (.docx) table or an Excel (.xlsx) sheet.

The header row is found by its column titles (Khmer or English), so the files
this system exports can be imported again, and so can a hand-made list with the
same columns in any order. Nothing is guessed: a cell that can't be read is
reported against its row and that row is skipped.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from app.core.errors import ValidationError

MAX_ROWS = 2000
_HEADER_SCAN_ROWS = 15

# field -> header fragments (matched against the header with spaces removed).
# Checked in this order, so "សាកលវិទ្យាល័យ" (university) is never taken for
# "វិទ្យាល័យ" (high school) and "លំដាប់ពិន្ទុ" (rank) never for "ថ្នាក់" (stream).
_HEADERS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("score_rank", ("លំដាប់", "ចំណាត់ថ្នាក់", "rank")),
    ("university", ("សាកល", "វិទ្យាស្ថាន", "university", "institute")),
    ("high_school", ("វិទ្យាល័យ", "វិទាល័យ", "highschool", "school")),
    ("note", ("ផ្សេងៗ", "សម្គាល់", "note", "remark")),
    ("phone", ("ទូរស័ព្ទ", "ទូរសព្ទ", "phone", "tel")),
    ("major", ("ជំនាញ", "មុខវិជ្ជា", "major")),
    ("gender", ("ភេទ", "gender", "sex")),
    ("grade", ("និទ្ទេស", "និទេស", "និ.", "grade")),
    ("full_name", ("គោត្តនាម", "គោតានាម", "ឈ្មោះ", "fullname", "name")),
    ("stream", ("ថ្នាក់", "ផ្នែក", "stream", "section")),
)

_KHMER_DIGITS = str.maketrans("០១២៣៤៥៦៧៨៩", "0123456789")
_EMPTY = {"", "-", "–", "—", "n/a", "na", "null", "none"}


@dataclass(slots=True)
class ParsedRow:
    row: int  # 1-based row number in the source file, for error messages
    values: dict[str, Any]


@dataclass(slots=True)
class ParsedList:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[tuple[int, str]] = field(default_factory=list)


def _header_field(text: str) -> str | None:
    key = re.sub(r"\s+", "", text).lower()
    if not key:
        return None
    for name, fragments in _HEADERS:
        if any(f in key for f in fragments):
            return name
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _gender(text: str) -> str | None:
    t = text.lower()
    if "ស្រី" in t or t in ("f", "female", "woman"):
        return "F"
    if "ប្រុស" in t or t in ("m", "male", "man"):
        return "M"
    return None


def _stream(text: str) -> str | None:
    t = text.lower()
    if "សង្គម" in t or "social" in t:
        return "social_science"
    if "វិទ្យាសាស្ត្រ" in t or "science" in t:
        return "science"
    return None


def _grade(text: str) -> str | None:
    m = re.fullmatch(r"([A-Ea-e])\.?", text.strip())
    return m.group(1).upper() if m else None


def _rank(text: str) -> int | None:
    m = re.fullmatch(r"(\d+)(?:\.0+)?", text.translate(_KHMER_DIGITS).strip())
    return int(m.group(1)) if m and int(m.group(1)) > 0 else None


def _phone(text: str, raw: Any) -> str:
    text = text.translate(_KHMER_DIGITS)
    # Excel stores 012 345 678 as the number 12345678 and drops the leading 0.
    if isinstance(raw, (int, float)) and re.fullmatch(r"[1-9]\d{7,8}", text):
        text = "0" + text
    return text


def _convert(rows: list[list[Any]], header_at: int, columns: dict[int, str]) -> ParsedList:
    out = ParsedList()
    for offset, cells in enumerate(rows[header_at + 1:], start=header_at + 2):
        raw = {fld: (cells[i] if i < len(cells) else None) for i, fld in columns.items()}
        text = {fld: _cell_text(v) for fld, v in raw.items()}
        text = {fld: ("" if v.lower() in _EMPTY else v) for fld, v in text.items()}
        if not any(text.values()):
            continue  # blank line or signature block
        if len(out.rows) + len(out.errors) >= MAX_ROWS:
            out.errors.append((offset, f"Only the first {MAX_ROWS} rows are imported."))
            break

        values: dict[str, Any] = {}
        problems: list[str] = []
        if not text.get("full_name"):
            problems.append("missing name")
        else:
            values["full_name"] = text["full_name"]

        if text.get("gender"):
            g = _gender(text["gender"])
            if g is None:
                problems.append(f"unknown gender “{text['gender']}”")
            else:
                values["gender"] = g
        else:
            problems.append("missing gender")

        if text.get("grade"):
            g = _grade(text["grade"])
            if g is None:
                problems.append(f"unknown grade “{text['grade']}” (use A–E)")
            else:
                values["grade"] = g
        if text.get("score_rank"):
            r = _rank(text["score_rank"])
            if r is None:
                problems.append(f"rank “{text['score_rank']}” is not a positive number")
            else:
                values["score_rank"] = r
        if text.get("stream"):
            s = _stream(text["stream"])
            if s is None:
                problems.append(f"unknown stream “{text['stream']}”")
            else:
                values["stream"] = s
        if text.get("phone"):
            values["phone"] = _phone(text["phone"], raw.get("phone"))
        for key in ("high_school", "university", "major", "note"):
            if text.get(key):
                values[key] = text[key]

        if problems:
            out.errors.append((offset, "; ".join(problems)))
        else:
            out.rows.append(ParsedRow(offset, values))
    return out


def _parse_grid(grid: list[list[Any]]) -> ParsedList | None:
    """Find the header row in a grid of cells and convert the rows below it."""
    for i, cells in enumerate(grid[:_HEADER_SCAN_ROWS]):
        columns: dict[int, str] = {}
        for col, cell in enumerate(cells):
            fld = _header_field(_cell_text(cell))
            if fld and fld not in columns.values():
                columns[col] = fld
        if "full_name" in columns.values() and len(columns) >= 2:
            return _convert(grid, i, columns)
    return None


def _xlsx_grids(data: bytes) -> list[list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [[list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()


def _docx_grids(data: bytes) -> list[list[list[Any]]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]


def parse_student_list(data: bytes) -> ParsedList:
    """Parse an .xlsx or .docx upload; raises ValidationError when it isn't a usable list."""
    try:
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile:
        raise ValidationError("Upload a Word (.docx) or Excel (.xlsx) file.") from None

    try:
        if any(n.startswith("xl/") for n in names):
            grids = _xlsx_grids(data)
        elif any(n.startswith("word/") for n in names):
            grids = _docx_grids(data)
        else:
            raise ValidationError("Upload a Word (.docx) or Excel (.xlsx) file.")
    except ValidationError:
        raise
    except Exception:
        raise ValidationError("This file could not be read. Is it a valid .docx or .xlsx file?") from None

    for grid in grids:
        parsed = _parse_grid(grid)
        if parsed is not None:
            return parsed
    raise ValidationError(
        "No student table found. The first rows need column titles such as "
        "គោត្តនាម-នាម and ភេទ (Name and Gender)."
    )
