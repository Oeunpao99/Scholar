"""Export students list to Word (.docx), Excel (.xlsx), and PDF (.pdf)."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from app.models.student import Student

KHMER_GENDER = {"M": "ប្រុស", "F": "ស្រី"}
KHMER_STREAM = {"science": "វិទ្យាសាស្ត្រ", "social_science": "វិទ្យាសាស្ត្រសង្គម"}


# ---------------------------------------------------------------------------
# Word (.docx) Export
# ---------------------------------------------------------------------------

def export_students_docx(
    students: list[Student],
    academic_year: int | None = None,
    institution: str | None = None,
    title: str | None = None,
) -> bytes:
    """Generate an official, beautifully styled Microsoft Word (.docx) document."""
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()

    # Page setup: A4 Landscape with 0.5 in margins for wide table readability
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    font_name = "Khmer OS Battambang"

    def apply_font(
        run,
        size_pt: float = 10,
        bold: bool = False,
        italic: bool = False,
        color_rgb: tuple[int, int, int] = (15, 23, 42),
    ) -> None:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        rPr = run._element.get_or_add_rPr()
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
        )
        rPr.append(rFonts)

    def set_cell_shading(cell, color_hex: str) -> None:
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tcPr.append(shd)

    def set_cell_padding(cell, top=100, bottom=100, left=120, right=120) -> None:
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    # 1. Header: Cambodian Official Header Layout
    # Left: Institution / Ministry; Right: Kingdom motto
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(5.3)
    header_table.columns[1].width = Inches(5.3)

    # Left cell: Institution
    left_cell = header_table.cell(0, 0)
    p_inst = left_cell.paragraphs[0]
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run(institution or "ប្រព័ន្ធគ្រប់គ្រងអាហារូបករណ៍")
    apply_font(r_inst, size_pt=11, bold=True, color_rgb=(30, 58, 138))

    # Right cell: Kingdom Motto
    right_cell = header_table.cell(0, 1)
    p_motto1 = right_cell.paragraphs[0]
    p_motto1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_motto1.paragraph_format.space_before = Pt(0)
    p_motto1.paragraph_format.space_after = Pt(2)
    r_m1 = p_motto1.add_run("ព្រះរាជាណាចក្រកម្ពុជា")
    apply_font(r_m1, size_pt=11, bold=True)

    p_motto2 = right_cell.add_paragraph()
    p_motto2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_motto2.paragraph_format.space_before = Pt(0)
    p_motto2.paragraph_format.space_after = Pt(2)
    r_m2 = p_motto2.add_run("ជាតិ សាសនា ព្រះមហាក្សត្រ")
    apply_font(r_m2, size_pt=10, bold=True)

    # 2. Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(2)
    main_title = title or "បញ្ជីរាយនាមសិស្សអាហារូបករណ៍"
    r_title = p_title.add_run(main_title)
    apply_font(r_title, size_pt=14, bold=True, color_rgb=(30, 58, 138))

    # Subtitle with academic year and totals
    total_count = len(students)
    female_count = sum(1 for s in students if s.gender == "F")
    male_count = sum(1 for s in students if s.gender == "M")

    year_str = f"ឆ្នាំសិក្សា {academic_year}" if academic_year else ""
    summary_str = f"សរុប៖ {total_count} នាក់ (ស្រី៖ {female_count} នាក់ · ប្រុស៖ {male_count} នាក់)"

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    sub_text = f"{year_str}  —  {summary_str}" if year_str else summary_str
    r_sub = p_sub.add_run(sub_text)
    apply_font(r_sub, size_pt=9.5, italic=True, color_rgb=(71, 85, 105))

    # 3. Data Table Setup
    # Columns definition: (Khmer Title, Width in inches, Alignment)
    headers = [
        ("ល.រ", 0.45, WD_ALIGN_PARAGRAPH.CENTER),
        ("គោត្តនាម-នាម", 1.45, WD_ALIGN_PARAGRAPH.LEFT),
        ("ភេទ", 0.50, WD_ALIGN_PARAGRAPH.CENTER),
        ("និទ្ទេស", 0.55, WD_ALIGN_PARAGRAPH.CENTER),
        ("លំដាប់ពិន្ទុ", 0.70, WD_ALIGN_PARAGRAPH.CENTER),
        ("វិទ្យាល័យ", 1.45, WD_ALIGN_PARAGRAPH.LEFT),
        ("ថ្នាក់", 0.95, WD_ALIGN_PARAGRAPH.CENTER),
        ("សាកលវិទ្យាល័យ/វិទ្យាស្ថាន", 1.65, WD_ALIGN_PARAGRAPH.LEFT),
        ("ជំនាញ/មុខវិជ្ជា", 1.25, WD_ALIGN_PARAGRAPH.LEFT),
        ("លេខទូរស័ព្ទ", 0.95, WD_ALIGN_PARAGRAPH.CENTER),
        ("ផ្សេងៗ", 0.80, WD_ALIGN_PARAGRAPH.LEFT),
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Apply table borders (subtle clean gray borders)
    tblPr = table._element.xpath("w:tblPr")
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_tr = table.rows[0]._element.get_or_add_trPr()
    hdr_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for col_idx, (col_name, col_width, col_align) in enumerate(headers):
        hdr_cells[col_idx].width = Inches(col_width)
        hdr_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(hdr_cells[col_idx], "1E293B")  # Slate 800 dark header
        set_cell_padding(hdr_cells[col_idx], top=120, bottom=120, left=80, right=80)

        p = hdr_cells[col_idx].paragraphs[0]
        p.alignment = col_align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(col_name)
        apply_font(r, size_pt=9, bold=True, color_rgb=(255, 255, 255))

    # Data Rows
    for row_idx, s in enumerate(students, start=1):
        row = table.add_row()
        trPr = row._element.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # Zebra striping: alternate row shading
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"

        values = [
            str(row_idx),
            s.full_name or "—",
            KHMER_GENDER.get(s.gender, s.gender or "—"),
            s.grade or "—",
            str(s.score_rank) if s.score_rank else "—",
            s.high_school or "—",
            KHMER_STREAM.get(s.stream, s.stream or "—") if s.stream else "—",
            s.university or "—",
            s.major or "—",
            s.phone or "—",
            s.note or "—",
        ]

        for col_idx, val in enumerate(values):
            cell = row.cells[col_idx]
            cell.width = Inches(headers[col_idx][1])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, bg_color)
            set_cell_padding(cell, top=80, bottom=80, left=80, right=80)

            p = cell.paragraphs[0]
            p.alignment = headers[col_idx][2]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            is_name = col_idx == 1
            apply_font(r, size_pt=8.5, bold=is_name, color_rgb=(15, 23, 42))

    # 4. Signature / Sign-off block
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(20)
    p_sign.paragraph_format.space_after = Pt(2)
    r_date = p_sign.add_run("ថ្ងៃ................ ទី........ ខែ........... ឆ្នាំ២០២...\n")
    apply_font(r_date, size_pt=9.5, italic=True)

    r_sig = p_sign.add_run("អ្នករៀបចំបញ្ជី\n\n\n\n")
    apply_font(r_sig, size_pt=10, bold=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel (.xlsx) Export
# ---------------------------------------------------------------------------

def export_students_xlsx(
    students: list[Student],
    academic_year: int | None = None,
    institution: str | None = None,
    title: str | None = None,
) -> bytes:
    """Generate an Excel (.xlsx) spreadsheet for the student list."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Student List"

    font_family = "Khmer OS Battambang"

    # Title block
    ws["A1"] = institution or "ប្រព័ន្ធគ្រប់គ្រងអាហារូបករណ៍"
    ws["A1"].font = Font(name=font_family, size=11, bold=True, color="1E3A8A")

    main_title = title or "បញ្ជីរាយនាមសិស្សអាហារូបករណ៍"
    if academic_year:
        main_title += f" (ឆ្នាំសិក្សា {academic_year})"
    ws["A2"] = main_title
    ws["A2"].font = Font(name=font_family, size=14, bold=True, color="0F172A")

    total_count = len(students)
    female_count = sum(1 for s in students if s.gender == "F")
    male_count = sum(1 for s in students if s.gender == "M")
    ws["A3"] = f"សរុប៖ {total_count} នាក់ (ស្រី៖ {female_count} នាក់ · ប្រុស៖ {male_count} នាក់)"
    ws["A3"].font = Font(name=font_family, size=10, italic=True, color="64748B")

    headers = [
        ("ល.រ", 8, "center"),
        ("គោត្តនាម-នាម", 22, "left"),
        ("ភេទ", 10, "center"),
        ("និទ្ទេស", 10, "center"),
        ("លំដាប់ពិន្ទុ", 14, "center"),
        ("វិទ្យាល័យ", 25, "left"),
        ("ថ្នាក់", 18, "center"),
        ("សាកលវិទ្យាល័យ/វិទ្យាស្ថាន", 30, "left"),
        ("ជំនាញ/មុខវិជ្ជា", 22, "left"),
        ("លេខទូរស័ព្ទ", 16, "center"),
        ("ផ្សេងៗ", 20, "left"),
    ]

    start_row = 5
    hdr_fill = PatternFill("solid", fgColor="1E293B")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx, (hdr_name, width, align) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=hdr_name)
        cell.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[start_row].height = 24

    for row_idx, s in enumerate(students, start=1):
        curr_row = start_row + row_idx
        ws.row_dimensions[curr_row].height = 20
        row_bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        cell_fill = PatternFill("solid", fgColor=row_bg)

        row_values = [
            row_idx,
            s.full_name or "—",
            KHMER_GENDER.get(s.gender, s.gender or "—"),
            s.grade or "—",
            s.score_rank or "—",
            s.high_school or "—",
            KHMER_STREAM.get(s.stream, s.stream or "—") if s.stream else "—",
            s.university or "—",
            s.major or "—",
            s.phone or "—",
            s.note or "—",
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            is_name = col_idx == 2
            cell.font = Font(name=font_family, size=9.5, bold=is_name)
            cell.fill = cell_fill
            cell.alignment = Alignment(horizontal=headers[col_idx - 1][2], vertical="center")
            cell.border = thin_border

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + max(len(students), 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF Export
# ---------------------------------------------------------------------------

def export_students_pdf(
    students: list[Student],
    academic_year: int | None = None,
    institution: str | None = None,
    title: str | None = None,
) -> bytes:
    """Generate a printable landscape PDF for the student list."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.services.export_service import ExportService

    # Use existing font registration from ExportService
    font_name = ExportService._register_font()
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "StudentTitle", parent=styles["Title"], fontName=font_name, fontSize=14, leading=18, alignment=1
    )
    sub_style = ParagraphStyle(
        "StudentSub", parent=styles["Normal"], fontName=font_name, fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#64748B")
    )
    hdr_style = ParagraphStyle(
        "StudentHdr", fontName=font_name, fontSize=8, leading=10, alignment=1, textColor=colors.white
    )
    cell_style = ParagraphStyle(
        "StudentCell", fontName=font_name, fontSize=7.5, leading=9.5
    )
    cell_center = ParagraphStyle(
        "StudentCellCenter", fontName=font_name, fontSize=7.5, leading=9.5, alignment=1
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    story: list = []
    main_title = title or "បញ្ជីរាយនាមសិស្សអាហារូបករណ៍"
    if academic_year:
        main_title += f" (ឆ្នាំសិក្សា {academic_year})"
    story.append(Paragraph(main_title, title_style))

    total_count = len(students)
    female_count = sum(1 for s in students if s.gender == "F")
    male_count = sum(1 for s in students if s.gender == "M")
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"សរុប៖ {total_count} នាក់ (ស្រី៖ {female_count} នាក់ · ប្រុស៖ {male_count} នាក់)", sub_style))
    story.append(Spacer(1, 4 * mm))

    col_widths = [10 * mm, 32 * mm, 12 * mm, 12 * mm, 16 * mm, 34 * mm, 24 * mm, 40 * mm, 32 * mm, 28 * mm, 26 * mm]
    headers = ["ល.រ", "គោត្តនាម-នាម", "ភេទ", "និទ្ទេស", "លំដាប់", "វិទ្យាល័យ", "ថ្នាក់", "សាកលវិទ្យាល័យ", "ជំនាញ", "ទូរស័ព្ទ", "ផ្សេងៗ"]

    table_data = [[Paragraph(h, hdr_style) for h in headers]]
    for idx, s in enumerate(students, start=1):
        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(s.full_name or "—", cell_style),
            Paragraph(KHMER_GENDER.get(s.gender, s.gender or "—"), cell_center),
            Paragraph(s.grade or "—", cell_center),
            Paragraph(str(s.score_rank) if s.score_rank else "—", cell_center),
            Paragraph(s.high_school or "—", cell_style),
            Paragraph(KHMER_STREAM.get(s.stream, s.stream or "—") if s.stream else "—", cell_center),
            Paragraph(s.university or "—", cell_style),
            Paragraph(s.major or "—", cell_style),
            Paragraph(s.phone or "—", cell_center),
            Paragraph(s.note or "—", cell_style),
        ])

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)

    doc.build(story)
    return buf.getvalue()
