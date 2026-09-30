"""Export services: CSV, Excel (xlsx) and PDF with Khmer font support."""

from __future__ import annotations

import csv
import io
import logging
from datetime import date
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.errors import ValidationError
from app.core.pagination import format_date_dmy, format_number
from app.models.user import User
from app.repositories.reference_repo import CategoryRepository
from app.repositories.report_repo import DailyReportRepository
from app.schemas.analytics import ExportRequest
from app.schemas.common import Counters
from app.services.audit_service import AuditService
from app.services.cumulative_service import GRADES, CumulativeService
from app.services.settings_service import SettingsService
from app.services.telegram_generator import (
    DEFAULT_GRADES,
    TelegramReportGenerator,
)

logger = logging.getLogger("scholar.export")

MIME_CSV = "text/csv; charset=utf-8"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_PDF = "application/pdf"

# Khmer-capable fonts shipped with the container image.
KHMER_FONTS = (
    "/usr/share/fonts/truetype/noto/NotoSansKhmer-Regular.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansKhmer-Regular.ttf",
    "C:/Windows/Fonts/Battambang.ttf",
    "C:/Windows/Fonts/KhmerUI.ttf",
)

HEADERS = {
    "date": "Date",
    "category": "Category",
    "grade": "Grade",
    "total": "Total",
    "female": "Female",
    "male": "Male",
    "pp": "PP",
    "kp": "KP",
    "cumulative_total": "Cumulative Total",
    "cumulative_female": "Cumulative Female",
    "cumulative_pp": "Cumulative PP",
    "cumulative_kp": "Cumulative KP",
}


class ExportService:
    def __init__(self, session: AsyncSession, current_user: User | None = None) -> None:
        self.session = session
        self.current_user = current_user
        self.repo = DailyReportRepository(session)
        self.category_repo = CategoryRepository(session)
        self.cumulative = CumulativeService(session)
        self.settings = SettingsService(session)
        self.audit = AuditService(session)

    # ------------------------------------------------------------- dispatch
    async def export(self, request: ExportRequest) -> tuple[bytes, str, str]:
        """Return ``(payload, filename, media_type)``."""
        start, end = self._resolve_period(request)
        year = await self.settings.get_current_year()
        rows = await self._collect(start, end, request)
        title = request.title or f"Daily Enrollment Report {format_date_dmy(start)} - {format_date_dmy(end)}"

        if request.format == "csv":
            payload = self._to_csv(rows, request)
            media, ext = MIME_CSV, "csv"
        elif request.format == "excel":
            payload = self._to_excel(rows, request, title, year)
            media, ext = MIME_XLSX, "xlsx"
        elif request.format == "pdf":
            payload = self._to_pdf(rows, request, title, year)
            media, ext = MIME_PDF, "pdf"
        else:
            raise ValidationError(f"Unsupported export format: {request.format}")

        filename = f"enrollment-report_{start.isoformat()}_{end.isoformat()}.{ext}"
        await self.audit.log(
            action="export",
            entity_type="export",
            summary=f"Exported {len(rows)} row(s) as {request.format.upper()} ({start} - {end})",
            changes={"format": request.format, "rows": len(rows)},
            user=self.current_user,
        )
        await self.session.commit()
        return payload, filename, media

    async def export_telegram_text(self, request: ExportRequest) -> tuple[bytes, str, str]:
        start, end = self._resolve_period(request)
        year = await self.settings.get_current_year()
        reports = await self.repo.list_in_range(start, end)
        by_date: dict[date, list] = {}
        for report in reports:
            by_date.setdefault(report.report_date, []).append(report)
        categories = await self.category_repo.list_ordered()
        generator = TelegramReportGenerator(zero_pad=request.zero_pad)

        chunks: list[str] = []
        for report_date in sorted(by_date):
            cumulative = await self.cumulative.cumulative_by_category(report_date)
            blocks = []
            for category in categories:
                report = next(
                    (r for r in by_date[report_date] if str(r.category_id) == str(category.id)), None
                )
                today_grades = (
                    {
                        entry.grade: Counters(
                            total=entry.today_total,
                            female=entry.today_female,
                            pp=entry.today_pp,
                            kp=entry.today_kp,
                        )
                        for entry in report.entries
                    }
                    if report
                    else {}
                )
                cum = next((c for c in cumulative if c.category_id == str(category.id)), None)
                blocks.append(
                    {
                        "roman_numeral": category.roman_numeral,
                        "title": category.render_title(year),
                        "today_grades": today_grades,
                        "cumulative": cum.counters() if cum else Counters(),
                    }
                )
            chunks.append(generator.render(report_date=report_date, categories=blocks).rstrip())

        text = f"សេចក្តីសង្ខេបប្រចាំឆ្នាំ {year}\n\n" + "\n\n\n".join(chunks) + "\n"
        return (
            text.encode("utf-8"),
            f"telegram-report_{start.isoformat()}_{end.isoformat()}.txt",
            "text/plain; charset=utf-8",
        )

    # ------------------------------------------------------------------ data
    async def _collect(
        self, start: date, end: date, request: ExportRequest
    ) -> list[dict[str, object]]:
        year = await self.settings.get_current_year()
        category_ids = [UUID(cid) for cid in request.category_ids] if request.category_ids else None
        reports = await self.repo.list_in_range(start, end, category_ids=category_ids)
        cumulative_cache: dict[tuple[date, str], dict] = {}

        rows: list[dict[str, object]] = []
        for report in sorted(reports, key=lambda r: (r.report_date, r.category.position)):
            key = (report.report_date, str(report.category_id))
            if request.include_cumulative and key not in cumulative_cache:
                cumulative_cache[key] = (
                    await self.cumulative.cumulative_for_category(
                        report.report_date, str(report.category_id)
                    )
                ).as_dict()
            cumulative = cumulative_cache.get(key, {})

            if request.include_grade_breakdown and report.entries:
                for entry in report.entries:
                    grade_cum = cumulative.get("grades", {}).get(entry.grade, {})
                    rows.append(
                        self._row(
                            report=report,
                            year=year,
                            grade=entry.grade,
                            total=entry.today_total,
                            female=entry.today_female,
                            pp=entry.today_pp,
                            kp=entry.today_kp,
                            cumulative=grade_cum,
                            include_cumulative=request.include_cumulative,
                        )
                    )
            else:
                rows.append(
                    self._row(
                        report=report,
                        year=year,
                        grade="",
                        total=report.today_total,
                        female=report.today_female,
                        pp=report.today_pp,
                        kp=report.today_kp,
                        cumulative={
                            "total": cumulative.get("total", 0),
                            "female": cumulative.get("female", 0),
                            "pp": cumulative.get("pp", 0),
                            "kp": cumulative.get("kp", 0),
                        },
                        include_cumulative=request.include_cumulative,
                    )
                )
        return rows

    @staticmethod
    def _row(
        *,
        report,
        year: int,
        grade: str,
        total: int,
        female: int,
        pp: int,
        kp: int,
        cumulative: dict,
        include_cumulative: bool,
    ) -> dict[str, object]:
        row: dict[str, object] = {
            "date": report.report_date.isoformat(),
            "date_dmy": format_date_dmy(report.report_date),
            "category_code": report.category.code,
            "category": report.category.render_title(year),
            "grade": grade,
            "total": total,
            "female": female,
            "male": max(total - female, 0),
            "pp": pp,
            "kp": kp,
        }
        if include_cumulative:
            row["cumulative_total"] = cumulative.get("total", 0)
            row["cumulative_female"] = cumulative.get("female", 0)
            row["cumulative_pp"] = cumulative.get("pp", 0)
            row["cumulative_kp"] = cumulative.get("kp", 0)
        return row

    @staticmethod
    def _resolve_period(request: ExportRequest) -> tuple[date, date]:
        if request.report_date:
            return request.report_date, request.report_date
        if request.start_date and request.end_date:
            if request.end_date < request.start_date:
                raise ValidationError("end_date must be on or after start_date.")
            return request.start_date, request.end_date
        if request.start_date:
            return request.start_date, date.today()
        today = date.today()
        return today.replace(day=1), today

    # ------------------------------------------------------------------ csv
    def _to_csv(self, rows: list[dict], request: ExportRequest) -> bytes:
        buffer = io.StringIO()
        # BOM keeps Excel happy with UTF-8 Khmer text.
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()) if rows else self._base_fields())
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {k: (format_number(v) if request.zero_pad and isinstance(v, int) else v)
                 for k, v in row.items()}
            )
        return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")

    @staticmethod
    def _base_fields() -> list[str]:
        return [
            "date", "date_dmy", "category_code", "category", "grade",
            "total", "female", "male", "pp", "kp",
            "cumulative_total", "cumulative_female", "cumulative_pp", "cumulative_kp",
        ]

    # ---------------------------------------------------------------- excel
    def _to_excel(
        self, rows: list[dict], request: ExportRequest, title: str, year: int
    ) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Daily Report"
        sheet["A1"] = title
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"] = f"Academic year: {year}"
        sheet["A2"].font = Font(italic=True, size=10)

        fields = [f for f in self._base_fields() if not rows or f in rows[0]]
        start_row = 4
        header_fill = PatternFill("solid", fgColor="0F172A")
        for column, field in enumerate(fields, start=1):
            cell = sheet.cell(row=start_row, column=column, value=HEADERS.get(field, field.title()))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for index, row in enumerate(rows, start=start_row + 1):
            for column, field in enumerate(fields, start=1):
                value = row.get(field, "")
                if request.zero_pad and isinstance(value, int):
                    value = format_number(value)
                sheet.cell(row=index, column=column, value=value)

        for column, field in enumerate(fields, start=1):
            width = max(len(field) + 2, 12)
            sheet.column_dimensions[get_column_letter(column)].width = (
                28 if field == "category" else width
            )
        sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)
        sheet.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(fields))}{start_row + max(len(rows), 1)}"
        )

        summary = workbook.create_sheet("Summary")
        summary["A1"] = "Totals"
        summary["A1"].font = Font(bold=True, size=13)
        for offset, metric in enumerate(("total", "female", "male", "pp", "kp"), start=3):
            summary.cell(row=offset, column=1, value=HEADERS[metric]).font = Font(bold=True)
            summary.cell(row=offset, column=2, value=sum(int(r.get(metric, 0)) for r in rows))
        summary.column_dimensions["A"].width = 18
        summary.column_dimensions["B"].width = 14

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ pdf
    def _to_pdf(self, rows: list[dict], request: ExportRequest, title: str, year: int) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "TitleKh", parent=styles["Title"], fontName=self._register_font(), fontSize=16, leading=20
        )
        cell_style = ParagraphStyle("CellKh", fontName=self._register_font(), fontSize=8, leading=10)

        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            title=title,
            author="Scholar Enrollment System",
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
        )
        story: list = [
            Paragraph(title, title_style),
            Paragraph(f"ឆ្នាំសិក្សា៖ {year}", cell_style),
            Spacer(1, 6 * mm),
        ]

        fields = [f for f in self._base_fields() if not rows or f in rows[0]]
        table_data = [[Paragraph(HEADERS.get(f, f.title()), cell_style) for f in fields]]
        for row in rows[:500]:
            table_data.append(
                [
                    Paragraph(
                        format_number(row.get(f)) if request.zero_pad and isinstance(row.get(f), int)
                        else str(row.get(f, "")),
                        cell_style,
                    )
                    for f in fields
                ]
            )

        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, -1), self._register_font()),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        if len(rows) > 500:
            story.append(Spacer(1, 4 * mm))
            story.append(
                Paragraph(
                    f"បង្ហាញតែ {500} ជួរដំបូងក្នុងសម្រាប់ការបង្ហាញ សូមប្រើ Excel សម្រាប់ទិន្នន័យពេញលេញ.",
                    cell_style,
                )
            )
        document.build(story)
        return buffer.getvalue()

    @staticmethod
    def _register_font() -> str:
        """Register a Khmer-capable font once; fall back to Helvetica."""
        from pathlib import Path

        for candidate in KHMER_FONTS:
            path = Path(candidate)
            if not path.exists():
                continue
            name = "ScholarKhmer"
            if name not in pdfmetrics.getRegisteredFontNames():
                try:
                    pdfmetrics.registerFont(TTFont(name, str(path)))
                except Exception:  # noqa: BLE001 - broken font file
                    logger.warning("Could not register Khmer font %s", path)
                    break
                return name
            return name
        logger.warning("No Khmer font found; PDF will use Helvetica.")
        return "Helvetica"
